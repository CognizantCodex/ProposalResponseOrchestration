from __future__ import annotations

from pathlib import Path


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


def extract_text(path: Path) -> tuple[str, int]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}. Supported: {sorted(SUPPORTED_EXTENSIONS)}")
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8-sig"), 1
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        return "\n\n".join(page.extract_text() or "" for page in reader.pages), len(reader.pages)
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            blocks.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "\n".join(blocks), max(1, len(document.sections))
    from pptx import Presentation

    deck = Presentation(path)
    blocks = [shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
    return "\n".join(blocks), len(deck.slides)


