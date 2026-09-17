from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

DEFAULT_CATEGORIES = "Application modernization; cloud migration; QEA automation; AI and analytics; infrastructure; cybersecurity; ERP; CRM; process automation."


def read_optional_text(path: Path, fallback: str) -> str:
    return path.read_text(encoding="utf-8-sig") if path.is_file() else fallback


def read_sls_reference(markdown_path: Path, employee_master_path: Path) -> str:
    if markdown_path.is_file():
        return markdown_path.read_text(encoding="utf-8-sig")
    if not employee_master_path.is_file():
        return "No SLS reference supplied. Return UNASSIGNED owners."
    workbook = load_workbook(employee_master_path, read_only=True, data_only=True)
    try:
        lines: list[str] = []
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                continue
            lines.append(" | ".join(str(value or "") for value in headers))
            lines.extend(" | ".join(str(value or "") for value in row) for row in rows)
        return "\n".join(lines)
    finally:
        workbook.close()

