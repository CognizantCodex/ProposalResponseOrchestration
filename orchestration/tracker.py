from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

HEADERS = ["Run ID", "Account", "File Name", "Source Path", "SHA-256", "Received At", "Status", "Current Agent", "Duplicate Of", "Error"]


class ExcelTracker:
    def __init__(self, path: Path):
        self.path = path

    def _load(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.exists():
            return load_workbook(self.path)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "RFP Progress"
        sheet.append(HEADERS)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:J1"
        workbook.save(self.path)
        return workbook

    def find_by_hash(self, digest: str) -> str | None:
        workbook = self._load()
        sheet = workbook["RFP Progress"]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[4] == digest and row[6] != "FAILED":
                return str(row[0])
        return None

    def insert(self, values: dict[str, Any]) -> None:
        workbook = self._load()
        sheet = workbook["RFP Progress"]
        sheet.append([values.get(header, "") for header in HEADERS])
        workbook.save(self.path)

    def update(self, run_id: str, *, status: str, agent: str, error: str = "") -> None:
        workbook = self._load()
        sheet = workbook["RFP Progress"]
        for row in sheet.iter_rows(min_row=2):
            if str(row[0].value) == run_id:
                row[6].value, row[7].value, row[9].value = status, agent, error
                workbook.save(self.path)
                return


