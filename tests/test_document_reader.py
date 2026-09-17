from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document

from orchestration.document_reader import extract_text


class DocumentReaderTest(TestCase):
    def test_text_and_markdown_extract_with_utf8_bom(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            for suffix in (".txt", ".md"):
                path = root / f"input{suffix}"
                path.write_text("\ufeffRequirement heading", encoding="utf-8")
                text, units = extract_text(path)
                self.assertEqual("Requirement heading", text)
                self.assertEqual(1, units)

    def test_docx_paragraphs_and_tables_extract(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "input.docx"
            document = Document()
            document.add_paragraph("Oracle Java migration")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Owner"
            table.rows[0].cells[1].text = "SEG"
            document.save(path)
            text, units = extract_text(path)
            self.assertIn("Oracle Java migration", text)
            self.assertIn("Owner | SEG", text)
            self.assertGreaterEqual(units, 1)

    def test_unsupported_extension_is_rejected(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "input.xls"
            path.write_text("not supported", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "Unsupported file type"):
                extract_text(path)

