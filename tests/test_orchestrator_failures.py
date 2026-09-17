from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from openpyxl import load_workbook

from orchestration.config import Settings
from orchestration.orchestrator import AgentOrchestrator


class AlwaysFailLlm:
    def generate(self, **kwargs):
        raise RuntimeError("model unavailable")


class OrchestratorFailureTest(TestCase):
    def test_retry_exhaustion_persists_failure_and_tracker_status(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "rfp.md"
            source.write_text("A valid requirement", encoding="utf-8")
            settings = Settings(tracker_path=root / "tracker.xlsx", state_dir=root / "runs", max_agent_retries=1)
            with self.assertRaisesRegex(RuntimeError, "model unavailable"):
                AgentOrchestrator(settings, llm=AlwaysFailLlm()).run(account="WellsFargo", source_path=source, run_id="failed-run")
            state = (root / "runs" / "failed-run.json").read_text(encoding="utf-8")
            self.assertIn('"status": "FAILED"', state)
            self.assertIn('"classifier": 2', state)
            workbook = load_workbook(root / "tracker.xlsx", read_only=True)
            try:
                sheet = workbook["RFP Progress"]
                row = list(sheet.iter_rows(min_row=2, values_only=True))[0]
                self.assertEqual("FAILED", row[6])
                self.assertEqual("classifier", row[7])
            finally:
                workbook.close()

    def test_validation_failure_is_recorded_without_tracker_row(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "rfp.exe"
            source.write_bytes(b"invalid")
            settings = Settings(tracker_path=root / "tracker.xlsx", state_dir=root / "runs", max_agent_retries=0)
            with self.assertRaises(ValueError):
                AgentOrchestrator(settings, llm=AlwaysFailLlm()).run(account="JPMC", source_path=source, run_id="invalid-run")
            state = (root / "runs" / "invalid-run.json").read_text(encoding="utf-8")
            self.assertIn('"status": "FAILED"', state)
            workbook = load_workbook(root / "tracker.xlsx", read_only=True)
            try:
                self.assertEqual(1, workbook["RFP Progress"].max_row)
            finally:
                workbook.close()

