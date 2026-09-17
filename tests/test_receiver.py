from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from orchestration.receiver_agent import DuplicateRfpError, ReceiverAgent
from orchestration.tracker import ExcelTracker


class ReceiverAgentTest(TestCase):
    def test_rejects_missing_and_unsupported_files(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            receiver = ReceiverAgent(ExcelTracker(root / "tracker.xlsx"), root / "Customer RFP Documentation")
            with self.assertRaises(FileNotFoundError):
                receiver.run(run_id="missing", account="Bank 1", source_path=root / "missing.docx")
            invalid = root / "input.exe"
            invalid.write_bytes(b"not an RFP")
            with self.assertRaises(ValueError):
                receiver.run(run_id="invalid", account="Bank 1", source_path=invalid)

    def test_empty_document_is_not_registered(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "empty.md"
            source.write_text("   ", encoding="utf-8")
            tracker = ExcelTracker(root / "tracker.xlsx")
            with self.assertRaises(ValueError):
                ReceiverAgent(tracker).run(run_id="empty", account="WellsFargo", source_path=source)
            self.assertFalse((root / "tracker.xlsx").exists())

    def test_github_blob_url_is_downloaded_and_registered(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            tracker = ExcelTracker(root / "tracker.xlsx")
            response = Mock()
            response.__enter__ = Mock(return_value=response)
            response.__exit__ = Mock(return_value=False)
            response.read.return_value = b"RFP requirements: migrate Oracle Java workloads."
            url = "https://github.com/CognizantCodex/ProposalResponseOrchestration/blob/develop/Customer%20RFP%20Documentation/Bank%201/Bank%201%20-%20Oracle%20Java%20Migration%20RFP%202026.md"
            with patch("orchestration.receiver_agent.urlopen", return_value=response) as download:
                metadata, text = ReceiverAgent(tracker, root / "Customer RFP Documentation").run(run_id="remote-1", account="Bank 1", source_path=url)
            download.assert_called_once()
            self.assertEqual("Bank 1 - Oracle Java Migration RFP 2026.md", metadata.file_name)
            self.assertIn("Oracle Java", text)
            self.assertTrue((root / "Customer RFP Documentation" / "Bank 1" / "Bank 1 - Oracle Java Migration RFP 2026" / "Response").is_dir())

    def test_duplicate_does_not_append_or_recreate(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "proposal.md"
            source.write_text("Cloud migration requirements", encoding="utf-8")
            tracker_path = root / "tracker.xlsx"
            receiver = ReceiverAgent(ExcelTracker(tracker_path), root / "Customer RFP Documentation")
            receiver.run(run_id="first", account="Bank 1", source_path=source)
            with self.assertRaises(DuplicateRfpError):
                receiver.run(run_id="second", account="Bank 1", source_path=source)
            workbook = load_workbook(tracker_path, read_only=True)
            try:
                self.assertEqual(2, workbook["RFP Progress"].max_row)
            finally:
                workbook.close()

    def test_github_url_conversion(self):
        converted = ReceiverAgent._raw_github_url("https://github.com/a/b/blob/develop/path/to/file.docx")
        self.assertEqual("https://raw.githubusercontent.com/a/b/develop/path/to/file.docx", converted)

